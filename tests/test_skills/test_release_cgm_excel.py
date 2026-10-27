import json
from pathlib import Path

from openpyxl import load_workbook

from deepyield.skills import get_skill, run_skill
from tests.test_compute.test_release_cgm_excel import report_fixture


def test_registered_skill_generates_cgm_workbook(tmp_path):
    source = tmp_path / "cgm.json"
    source.write_text(json.dumps(report_fixture()), encoding="utf-8")
    assert get_skill("release_cgm_excel").category == "reporting"
    result = run_skill("release_cgm_excel", {"file_paths": [str(source)],
                                             "params": {"output_dir": str(tmp_path)}})
    assert result.success, result.errors
    assert Path(result.artifacts["xlsx"]).is_file()
    assert load_workbook(result.artifacts["xlsx"]).sheetnames == ["raw"]
    assert result.data["sheet_rows"] == {"raw": 1}
    json.dumps(result.to_dict(serialize=True), allow_nan=False)


def test_registered_skill_validates_arguments():
    for payload in ({}, {"file_paths": ["a", "b"]}, {"file_paths": ["a"]},
                    {"file_paths": ["a"], "params": {"output_dir": "out", "threshold": 0}}):
        result = run_skill("release_cgm_excel", payload)
        assert not result.success and result.errors
