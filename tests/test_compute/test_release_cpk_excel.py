"""Workbook contract, exact thresholds, provenance and saved final remarks."""
from copy import deepcopy
from datetime import date
from io import BytesIO
import json

from openpyxl import load_workbook
import pytest

from deepyield.compute.release_cpk_excel import render_cpk_xlsx
from deepyield.compute.release_cpk_excel_input import export_release_cpk_xlsx
from deepyield.rules.loader import load_analysis_rules


def report_fixture():
    catalog = load_analysis_rules()
    conditions = {"part_type": "VP150BC", "test_temperature": "25", "test_code": "FT"}
    def item(nr, cpk, unit="V", **extra):
        return {"group_id": str(nr), "column": f"{nr}_TEST", "test_num": nr,
                "test_name": f"TEST_{nr}", "unit": unit, "n": 50, "fail_count": 0,
                "effective_lsl": 0, "effective_usl": 10, "mean": 5, "std": .1,
                "cpk_cp": .99, "cpk": cpk, "status": "PASS", "conditions": dict(conditions),
                **extra}
    items = [item(1, 2), item(2, 1.67), item(3, 1.66), item(4, 500),
             item(5, 500.01), item(6, .5, " "), item(7, 900, None),
             item(8, None, reason="ZERO_OR_UNAVAILABLE_STD", status="NOT_CALCULABLE"),
             item(9, float("inf"))]
    return {"report_type": "release_bin1_cpk", "items": items,
            "unexecuted_tests": [], "source": {"source_file": "synthetic.stdf"},
            "rules": {"catalog_hash": catalog.catalog_hash, "bands": [
                {"rule_id": f"CPK-00{i}", "condition": catalog.rule(f"CPK-00{i}").spec["condition"],
                 "status": catalog.rule(f"CPK-00{i}").status} for i in range(1, 4)]}}


def workbook(report=None, **kwargs):
    output = render_cpk_xlsx(report or report_fixture(), report_date=date(2026, 10, 19), **kwargs)
    return output, load_workbook(BytesIO(output["content"]))


def test_exact_four_sheets_columns_boundaries_colors_and_empty_high_remark():
    output, book = workbook()
    assert output["filename"] == "20261019_VP150BC_25_Bin1_CPK_Report.xlsx"
    assert book.sheetnames == ["raw", "low_CPK_Item", "CPK_too_Good", "Test_without_unit"]
    headers = [None, "Test Nr.", "Test Name", "Unit", "ExecCount", "FailCount",
               "Lowlim", "Uplim", "Mean", "Std", "cpk/cp", "cpk"]
    for index, sheet in enumerate(book):
        assert sheet.max_column == (12 if index == 0 else 13)
        assert [cell.value for cell in sheet[1]] == headers + (["Remark"] if index else [])
        assert sheet.freeze_panes == "E2"
        assert sheet.auto_filter.ref
    assert [row[1].value for row in book["low_CPK_Item"].iter_rows(min_row=2)] == [3, 2]
    assert [row[1].value for row in book["CPK_too_Good"].iter_rows(min_row=2)] == [5]
    assert book["CPK_too_Good"]["M2"].value is None
    assert [row[1].value for row in book["Test_without_unit"].iter_rows(min_row=2)] == [6, 7]
    assert book["raw"]["A2"].value == "000001 TEST_1 V"
    assert book["raw"]["G2"].value == 0
    assert [book["raw"][f"L{i}"].fill.fgColor.rgb[-6:] for i in (2, 3, 4)] == [
        "CCE9DF", "F9E8BB", "F4D2D4"]
    assert book["raw"]["L10"].value is None  # Infinity cannot become too good.
    assert book["raw"]["A2"].comment
    assert not book["low_CPK_Item"]["M2"].value  # No history, site difference or manual note.


def test_saved_page_final_remark_wins_and_high_is_always_blank():
    report = report_fixture()
    report["items"][2]["diagnosis"] = {"remark": "Old automatic remark"}
    report["items"][4]["diagnosis"] = {"remark": "Must never appear on high sheet"}
    result = {"rows": [{"key": "merged", "column": "3_TEST",
                       "cpk": {"group_id": "3", "diagnosis": {"remark": "Engineer final remark"}}}]}
    before = deepcopy(report)
    _, book = workbook(report, result=result)
    assert book["low_CPK_Item"]["M2"].value == "Engineer final remark"
    assert book["CPK_too_Good"]["M2"].value is None
    assert report == before


def test_saved_custom_bands_drive_filter_and_colors():
    report = report_fixture()
    report["rules"]["bands"][0]["condition"] = "< 1"
    report["rules"]["bands"][1]["condition"] = ">= 1 and < 1.5"
    report["rules"]["bands"][2]["condition"] = ">= 1.5"
    _, book = workbook(report)
    assert book["low_CPK_Item"].max_row == 1
    assert book["low_CPK_Item"].max_column == 13
    assert book["raw"]["L4"].fill.fgColor.rgb[-6:] == "CCE9DF"


def test_unexecuted_and_spec_variants_are_preserved():
    report = report_fixture()
    report["items"].append({**report["items"][0], "group_id": "another_spec", "effective_usl": 12})
    report["unexecuted_tests"] = [{"column": "missing", "n": 0, "status": "INSUFFICIENT_DATA",
                                  "reason": "NO_SELECTED_EXECUTION"}]
    metadata = {"test_catalog": [{"column": "missing", "test_num": 99, "name": "NOT_RUN",
                                 "unit": "V", "low_limit": 1, "high_limit": 2}]}
    output, book = workbook(report, metadata=metadata)
    assert output["sheet_rows"]["raw"] == 11
    assert book["raw"]["B12"].value == 99
    assert book["raw"]["E12"].value == 0
    assert book["raw"]["L12"].value is None
    assert book["raw"]["A12"].comment


def test_missing_or_mixed_metadata_is_optional_in_filenames():
    report = report_fixture()
    report["items"][0]["conditions"]["part_type"] = "OTHER"
    output, _ = workbook(report)
    assert output['filename'] == '20261019_25_Bin1_CPK_Report.xlsx'
    _, book = workbook(report, device_id="EXPLICIT")
    assert book.sheetnames[0] == "raw"
    for item in report["items"]:
        item["conditions"] = {}
    output, _ = workbook(report, device_id="D")
    assert output['filename'] == '20261019_D_Bin1_CPK_Report.xlsx'
    output, _ = workbook(report)
    assert output['filename'] == '20261019_Bin1_CPK_Report.xlsx'
    output, _ = workbook(report, device_id="D", test_temp="MSR_Release_EMTT")
    assert "MSR_Release_EMTT" in output["filename"]


def test_stage_fallback_and_safe_text_not_formula():
    report = report_fixture()
    for item in report["items"]:
        item["conditions"].pop("test_temperature")
    report["items"][2]["test_name"] = "=HYPERLINK(\"bad\")"
    report["items"][2]["diagnosis"] = {"remark": "=1+1"}
    output, book = workbook(report)
    assert "_FT_" in output["filename"]
    assert book["raw"]["C4"].data_type == "s"
    assert book["low_CPK_Item"]["M2"].data_type == "s"


def test_json_export_rejects_overwrite_and_new_selection(tmp_path):
    source = tmp_path / "run.json"
    report = report_fixture()
    report["items"] = report["items"][:8]
    source.write_text(json.dumps({"reports": {"cpk": report}, "result": {}}), encoding="utf-8")
    result = export_release_cpk_xlsx(source, tmp_path / "output")
    assert load_workbook(result["path"]).sheetnames[2] == "CPK_too_Good"
    with pytest.raises(FileExistsError):
        export_release_cpk_xlsx(source, tmp_path / "output")
    with pytest.raises(ValueError, match="固定"):
        export_release_cpk_xlsx(source, tmp_path / "other", sites=[1])
    with pytest.raises(ValueError, match="固定"):
        export_release_cpk_xlsx(source, tmp_path / "other", rule_dir="config/analysis_rule")


def test_parquet_route_uses_existing_bin1_calculation(monkeypatch, tmp_path):
    from tests.test_compute.test_release_cpk import dataset
    from deepyield.compute.release_cpk import build_release_bin1_cpk_report
    import deepyield.compute.release_cpk_excel_input as entry
    frame, meta = dataset()
    monkeypatch.setattr(entry, "read_release_wide_all", lambda path: (frame, meta))
    output = entry.export_release_cpk_xlsx(tmp_path / "fixture.parquet", tmp_path)
    book = load_workbook(output["path"])
    expected = build_release_bin1_cpk_report(frame, meta, catalog=load_analysis_rules())["items"][0]
    assert book["raw"]["E2"].value == expected["n"]
    assert book["raw"]["L2"].value == pytest.approx(expected["cpk"])
    assert book["raw"]["K2"].value == pytest.approx(expected["cpk_cp"])


def test_invalid_report_and_rules_rejected():
    with pytest.raises(ValueError, match="bin1"):
        workbook({"report_type": "other"})
    report = report_fixture()
    report["rules"]["bands"] = []
    with pytest.raises(ValueError, match="规则"):
        workbook(report)

def test_finite_but_insufficient_cpk_is_numeric_screen_only_with_gray_color():
    report = report_fixture()
    report["items"][2].update(status="INSUFFICIENT_DATA", rule_status=None,
                              reason="TOO_FEW_VALID_SAMPLES", n=2)
    _, book = workbook(report)
    assert book["low_CPK_Item"]["B2"].value == 3
    assert book["low_CPK_Item"]["L2"].fill.fgColor.rgb[-6:] == "E5EAED"
    provenance = json.loads(book["low_CPK_Item"]["A2"].comment.text)
    assert provenance["status"] == "INSUFFICIENT_DATA"
    assert provenance["numeric_screen_only"]
    assert "sample" in book["low_CPK_Item"]["M2"].value.lower()


def test_cross_source_unexecuted_identity_and_final_remark():
    report = report_fixture()
    report["unexecuted_tests"] = [{"column": "missing", "n": 0, "status": "INSUFFICIENT_DATA",
                                  "reason": "NO_SELECTED_EXECUTION"}]
    result = {"rows": [{"key": "unexecuted:sourceA:missing", "column": "sourceA:missing",
                       "number": 99, "name": "NOT_RUN", "unit": None,
                       "cpk": {"dataset_id": "sourceA", "group_id": None,
                               "diagnosis": {"remark": "Saved unexecuted final"}}}]}
    _, book = workbook(report, result=result)
    assert book["raw"]["B11"].value == 99
    sheet = book["Test_without_unit"]
    assert sheet["B4"].value == 99
    assert sheet["M4"].value == "Saved unexecuted final"


def test_bare_json_missing_unexecuted_identity_rejected_or_completed_with_catalog(tmp_path):
    report = report_fixture()
    report["items"] = report["items"][:8]
    report["unexecuted_tests"] = [{"column": "missing", "n": 0, "status": "INSUFFICIENT_DATA",
                                  "reason": "NO_SELECTED_EXECUTION"}]
    source = tmp_path / "missing.json"
    source.write_text(json.dumps(report), encoding="utf-8")
    with pytest.raises(ValueError, match="未执行测项"):
        export_release_cpk_xlsx(source, tmp_path / "out")
    metadata = {"test_catalog": [{"column": "missing", "test_num": 99, "name": "NOT_RUN", "unit": "V"}]}
    source.write_text(json.dumps({"reports": {"cpk": report}, "metadata": metadata}), encoding="utf-8")
    output = export_release_cpk_xlsx(source, tmp_path / "out")
    assert load_workbook(output["path"])["raw"]["B10"].value == 99
