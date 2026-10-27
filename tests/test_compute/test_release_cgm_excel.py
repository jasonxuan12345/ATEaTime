"""Exact Cgm workbook coordinates, selected statistics and joint color contract."""
from copy import deepcopy
from datetime import date
from io import BytesIO
import json

from openpyxl import load_workbook
import pytest

from deepyield.compute.release_cgm import build_release_cgm_report
from deepyield.compute.release_cgm_excel import render_cgm_xlsx
from deepyield.compute.release_cgm_excel_input import export_release_cgm_xlsx
from deepyield.rules.loader import load_analysis_rules
from tests.test_compute.test_release_cpk import dataset

COLORS = {"Pass": "74B230", "Need to check": "F0EA00", "Fail": "EA0000"}


def report_fixture(cgs=(1.33, 1.5), cpk=None):
    sites = [dict(site=index + 1, n=25, mean=5., std=1., min=3., max=7.,
                  cg=cg, fail_count=0) for index, cg in enumerate(cgs)]
    item = dict(group_id="g", column="T", test_num=1001, test_name="VDD", unit="V",
                effective_lsl=0., effective_usl=10., sites=sites,
                supported_by_bin1_cpk=cpk, conditions={"part_type": "DeviceA"})
    return dict(report_type="release_cgm", items=[item], sites=list(range(1, len(cgs) + 1)),
                unexecuted_tests=[], source={"source_file": "fixture.stdf"}, policy={"cg_tolerance_fraction": .20})


def workbook(report=None, **kwargs):
    result = render_cgm_xlsx(report if report is not None else report_fixture(),
                             report_date=date(2026, 10, 20), **kwargs)
    return result, load_workbook(BytesIO(result["content"]))


def color(cell):
    return cell.fill.fgColor.rgb[-6:] if cell.fill.patternType else None


def test_exact_filename_three_headers_and_two_site_coordinates():
    report = report_fixture()
    report["sites"] = [2, 1]
    report["items"][0]["sites"].reverse()
    before = deepcopy(report)
    result, book = workbook(report)
    assert report == before
    assert result["filename"] == "20261020_DeviceA_Cgm_Report.xlsx"
    assert result["sheet_rows"] == {"raw": 1}
    assert book.sheetnames == ["raw"]
    sheet = book["raw"]
    assert sheet.max_column == 23 and sheet.max_row == 4
    assert not sheet.merged_cells.ranges
    assert [cell.value for cell in sheet[1]] == [None] * 5 + ["Site"] + [None] * 6 + ["Site"] + [None] * 10
    assert [cell.value for cell in sheet[2]] == [None] * 5 + [1] + [None] * 6 + [2] + [None] * 10
    site_headers = ["Mean", "StdDev", "min", "max", "cg*", "ExecCount", "FailCount"]
    assert [cell.value for cell in sheet[3]] == ["Test Nr.", "Test Name", "Unit", "Lowlim", "Uplim"] + site_headers * 2 + ["min Cg", "Bin1 CPK", "Result", "Remark"]
    assert [cell.value for cell in sheet[4]][:12] == [1001, "VDD", "V", 0, 10, 5, 1, 3, 7, 1.33, 25, 0]
    assert sheet.freeze_panes == "F4" and sheet.auto_filter.ref == "A3:W4"
    assert len(sheet.conditional_formatting) == 5  # J, Q, T, U and V
    assert [color(sheet[address]) for address in ("J4", "Q4", "T4", "V4")] == [COLORS["Pass"]] * 4
    assert color(sheet['U4']) is None
    assert not sheet['W4'].value and sheet['W4'].comment  # Passing item needs no generic remark.
    assert sheet["U4"].value is None and sheet["U4"].has_style
    assert sheet["A4"].comment


@pytest.mark.parametrize("cg,cpk,result", [
    (1.33, None, "Pass"), (1.33, -1, "Pass"), (2, .5, "Pass"),
    (1.329999, None, "Need to check"), (1.0, 1.999999, "Need to check"),
    (1.0, 2, "Pass"), (1.2, 2.1, "Pass"), (1.2, float("nan"), "Need to check"),
    (.999999, None, "Fail"), (0, 1.999999, "Fail"), (-1, .5, "Fail"),
    (.99, 2, "Pass"), (.5, float("nan"), "Fail"),
])
def test_joint_thresholds_and_independent_cg_colors(cg, cpk, result):
    _, book = workbook(report_fixture((cg,), cpk))
    sheet = book["raw"]
    own = "Pass" if cg >= 1.33 else "Need to check" if cg >= 1 else "Fail"
    assert sheet["O4"].value == result
    assert color(sheet["O4"]) == COLORS[result]
    assert sheet["M4"].value == cg
    assert color(sheet["J4"]) == color(sheet["M4"]) == COLORS[own]


@pytest.mark.parametrize("equal", [False, True])
def test_zero_std_preserves_infinite_or_nan_minimum(equal):
    report = report_fixture()
    item = report["items"][0]
    if equal:
        item["effective_usl"] = item["effective_lsl"]
    for site in item["sites"]:
        site.update(std=0, cg=None)
    _, book = workbook(report)
    sheet = book["raw"]
    assert sheet["J4"].value == sheet["Q4"].value == ("NaN" if equal else "Infinity")
    assert sheet["G4"].value == sheet["N4"].value == 0
    assert color(sheet["J4"]) == (None if equal else COLORS["Pass"])
    assert sheet['T4'].value == ('NaN' if equal else 'Infinity')
    assert color(sheet['T4']) == (None if equal else COLORS['Pass'])
    assert sheet['V4'].value == ('Need to check' if equal else 'Pass')
    assert color(sheet['U4']) is None


@pytest.mark.parametrize('cgs,expected', [
    (('Infinity', 'Infinity'), 'Infinity'), ((float('inf'), 1.2), 1.2),
    (('NaN', 'NaN'), 'NaN'), (('NaN', 0), 'NaN'), ((1.5, 'NaN'), 'NaN'),
    (('Infinity', 'NaN'), 'NaN'), ((float('nan'), 1.5), 'NaN'),
    ((None, 1.5), 1.5), ((None, None), None),
])
def test_minimum_keeps_special_values_and_propagates_any_nan(cgs, expected):
    _, book = workbook(report_fixture(cgs))
    sheet = book['raw']
    assert sheet['T4'].value == expected
    if expected == 'NaN':
        assert color(sheet['T4']) is None and sheet['V4'].value == 'Need to check'


@pytest.mark.parametrize('cpk,expected', [(None, None), (float('nan'), None), (-1, 'Fail'),
    (1.669999, 'Fail'), (1.67, 'Need to check'), (1.999999, 'Need to check'), (2, 'Pass'), (2.01, 'Pass')])
@pytest.mark.parametrize('cg', [.5, 1.2, 1.5])
def test_bin1_cpk_has_its_own_numeric_colors(cpk, expected, cg):
    _, book = workbook(report_fixture((cg,), cpk))
    assert color(book['raw']['N4']) == COLORS.get(expected)


def test_final_saved_remark_wins_and_is_formula_safe_without_mutating_report():
    report = report_fixture((.5,))
    report['items'][0]['diagnosis'] = {'remark': 'Earlier conclusion'}
    result = {'rows': [{'column': 'different-source:T', 'cgm': {
        'group_id': 'g', 'dataset_id': 'different-source', 'diagnosis': {
            'remark': '=工程确认后的结论', 'knowledge_version': 'saved-version'}}}]}
    before = deepcopy(report)
    _, book = workbook(report, result=result)
    cell = book['raw']['P4']
    assert cell.value == '=工程确认后的结论' and cell.data_type == 's'
    assert json.loads(cell.comment.text)['knowledge_version'] == 'saved-version'
    assert report == before


def test_automatic_remark_uses_generic_guidance_and_synthetic_reference(monkeypatch):
    from tests.test_compute.remark_fixtures import example_catalog
    monkeypatch.setattr("deepyield.compute.release_remarks.load_knowledge", example_catalog)
    report = report_fixture((.5,))
    report['items'][0].update(test_num=1001, test_name='ExampleVoltage',
                              cg=.5, status='FAIL', cpk_override_applied=False)
    _, book = workbook(report)
    cell = book['raw']['P4']
    assert cell.value.startswith('1. Check site repeatability, contact and distribution.\n2. ')
    assert 'Synthetic reference: review the example limits.' in cell.value
    assert 'Min Cg=' not in cell.value and 'Historical project reference' not in cell.value
    diagnostic = json.loads(cell.comment.text)
    assert diagnostic['knowledge_hash'] and diagnostic['knowledge']['status'] == 'reference'


def test_infinity_does_not_mask_real_finite_zero_and_missing_site_is_blank():
    report = report_fixture((None, 0))
    report["sites"].append(3)
    report["items"][0]["sites"][0].update(std=0)
    _, book = workbook(report)
    sheet = book["raw"]
    assert sheet["J4"].value == "Infinity"
    assert sheet["Q4"].value == 0
    assert [sheet.cell(4, col).value for col in range(20, 27)] == [None] * 7
    assert sheet["AA4"].value == 0 and color(sheet["AA4"]) == COLORS["Fail"]
    assert sheet["AC4"].value == "Fail"


def test_unexecuted_rows_catalog_order_missing_unit_precision_and_formula_safe_text():
    report = report_fixture()
    item = report["items"][0]
    item.update(unit=None, effective_lsl=.123456789012345, test_name="=1+1")
    report["test_catalog"] = [dict(column="never", test_num=9000, name="NEVER", unit=None),
                               dict(column="T", test_num=1001, name="=1+1")]
    report["unexecuted_tests"] = [dict(column="never", n=0)]
    _, book = workbook(report)
    sheet = book["raw"]
    assert sheet["A4"].value == 9000 and sheet["A5"].value == 1001
    assert sheet["C4"].value is None and sheet["C5"].value is None
    assert sheet["D5"].value == .123456789012345
    assert sheet["B5"].value == "=1+1" and sheet["B5"].data_type == "s"
    assert all(sheet.cell(4, col).value is None for col in range(6, 20))
    assert sheet["T4"].value is None and color(sheet["T4"]) is None
    assert sheet["V4"].value == "Need to check"


@pytest.mark.parametrize('field,column', [('mean', 'F'), ('std', 'G'), ('min', 'H'), ('max', 'I'), ('cg', 'J'), ('n', 'K'), ('fail_count', 'L')])
def test_saved_result_with_missing_site_statistics_leaves_missing_cells_blank(field, column):
    report = report_fixture()
    del report['items'][0]['sites'][0][field]
    _, book = workbook(report)
    assert book['raw'][f'{column}4'].value is None


def test_saved_selection_and_json_export_no_overwrite_or_reselection(tmp_path):
    report = report_fixture((.9,), 2)
    report.update(selection={"sites": [1]}, bin1_cpk_source={"source_file": "other.stdf"})
    source = tmp_path / "run.json"
    source.write_text(json.dumps({"reports": {"cgm": report}}), encoding="utf-8")
    output = export_release_cgm_xlsx(source, tmp_path / "out")
    book = load_workbook(output["path"])
    assert book["raw"]["N4"].value == 2 and book["raw"]["O4"].value == "Pass"
    provenance = json.loads(book.properties.description)
    assert provenance["selection"] == {"sites": [1]}
    assert provenance["bin1_cpk_source"] == {"source_file": "other.stdf"}
    with pytest.raises(FileExistsError):
        export_release_cgm_xlsx(source, tmp_path / "out")
    for params in ({"sites": [2]}, {"part_indices": [1]}, {"rule_dir": "x"},
                   {"minimum_samples_per_site": 2}, {"bin1_cpk_report_path": "x"}):
        with pytest.raises(ValueError, match="固定"):
            export_release_cgm_xlsx(source, tmp_path / "other", **params)


def test_new_site_statistics_and_selected_parquet_route(monkeypatch, tmp_path):
    frame, meta = dataset((4, 5, 6, 7, 9, 11))
    frame["SITE"] = [1, 1, 1, 2, 2, 2]
    report = build_release_cgm_report(frame, meta, catalog=load_analysis_rules())
    first, second = report["items"][0]["sites"]
    assert (first["mean"], first["std"], first["min"], first["max"], first["fail_count"]) == (5, 1, 4, 6, 0)
    assert (second["mean"], second["std"], second["min"], second["max"], second["fail_count"]) == (9, 2, 7, 11, 1)
    import deepyield.compute.release_cgm_excel_input as entry
    monkeypatch.setattr(entry, "read_release_wide_all", lambda path: (frame, meta))
    output = export_release_cgm_xlsx(tmp_path / "fixture.parquet", tmp_path / "out",
                                      sites=[2], part_indices=[4, 5, 6])
    sheet = load_workbook(output["path"])["raw"]
    assert sheet["F2"].value == 2 and sheet.max_column == 16
    assert sheet["F4"].value == 9 and sheet["G4"].value == 2
    assert sheet["K4"].value == 3 and sheet["L4"].value == 1
    assert sheet["M4"].value == pytest.approx(.20 * 8 / 12)


def test_native_stdf_to_parquet_to_excel_exact_test_order_and_statistics(tmp_path):
    from tests.test_compute.test_release_stdf import convert, part, ptr, specification
    body = b""
    for site, values in ((2, (4., 5., 6.)), (1, (5., 5., 5.))):
        for index, value in enumerate(values):
            body += part(site=site, name=f"P{site}_{index}", tests=
                ptr(value, number=9000, site=site, name="FIRST", optional=specification(low=0, high=10)) +
                ptr(value, number=1001, site=site, name="SECOND", optional=specification(low=0, high=10)))
    (frame, meta), _, path = convert(tmp_path, body)
    assert [test["test_num"] for test in meta["test_catalog"]] == [9000, 1001]
    output = export_release_cgm_xlsx(path, tmp_path / "out", device_id="NativeTest")
    sheet = load_workbook(output["path"])["raw"]
    assert [sheet.cell(row, 1).value for row in (4, 5)] == [9000, 1001]
    assert (sheet["F2"].value, sheet["M2"].value) == (1, 2)
    assert sheet["F4"].value == 5 and sheet["G4"].value == 0 and sheet["J4"].value == "Infinity"
    assert sheet["M4"].value == 5 and sheet["N4"].value == 1
    assert sheet["T4"].value == pytest.approx(.20 * 10 / 6)


def test_invalid_input_missing_device_and_spec_variants():
    with pytest.raises(ValueError, match="CGM"):
        workbook({"report_type": "other"})
    report = report_fixture()
    report["items"].append({**deepcopy(report["items"][0]), "group_id": "g2", "effective_usl": 12})
    _, book = workbook(report)
    assert book["raw"].max_row == 5 and book["raw"]["E5"].value == 12
    for item in report["items"]:
        item["conditions"] = {}
    output, _ = workbook(report)
    assert output['filename'] == '20261020_Cgm_Report.xlsx'
    output, _ = workbook(report, metadata={"stdf_metadata": {"mir": {"part_type": "MIR_PRODUCT"}}})
    assert output["filename"] == "20261020_MIR_PRODUCT_Cgm_Report.xlsx"


def test_mixed_device_and_missing_test_descriptions_still_export():
    report = report_fixture()
    report['items'].append({**deepcopy(report['items'][0]), 'group_id': 'other',
                           'conditions': {'part_type': 'Other'}})
    report['items'][0].pop('test_name')
    report['items'][0].pop('test_num')
    report['unexecuted_tests'] = [{'column': 'unknown', 'n': 0}]
    output, book = workbook(report)
    assert output['filename'] == '20261020_Cgm_Report.xlsx'
    assert book['raw']['A4'].value is None and book['raw']['B4'].value is None
    assert book['raw'].max_row == 6


def test_conditional_format_formulas_cover_boundaries_blank_cpk_and_infinity():
    _, book = workbook()
    sheet = book["raw"]
    rules = {str(key.sqref): value for key, value in sheet.conditional_formatting._cf_rules.items()}
    assert [rule.formula[0] for rule in rules["J4"]] == [
        'OR(J4="Infinity",AND(ISNUMBER(J4),J4>=1.33))',
        'AND(ISNUMBER(J4),J4>=1.0,J4<1.33)', 'AND(ISNUMBER(J4),J4<1.0)']
    assert rules['T4'][0].formula[0] == 'OR(T4="Infinity",AND(ISNUMBER(T4),T4>=1.33))'
    assert [rule.formula[0] for rule in rules['U4']] == [
        'AND(ISNUMBER(U4),U4>=2.0)', 'AND(ISNUMBER(U4),U4>=1.67,U4<2.0)', 'AND(ISNUMBER(U4),U4<1.67)']
    assert [rule.formula[0] for rule in rules["V4"]] == [
        '$V4="Pass"', '$V4="Need to check"', '$V4="Fail"']
    assert [rule.dxf.fill.bgColor.rgb[-6:] for rule in rules["V4"]] == list(COLORS.values())


def test_single_sample_and_empty_site_keep_unavailable_statistics_blank():
    frame, metadata = dataset((5, 6, 7))
    frame["SITE"] = [0, 1, 1]
    frame.at[2, "__test_records"] = []
    report = build_release_cgm_report(frame, metadata, catalog=load_analysis_rules())
    _, book = workbook(report)
    sheet = book["raw"]
    assert sheet["F4"].value == 5 and sheet["G4"].value is None
    assert sheet["J4"].value is None and sheet["K4"].value == 1
    assert sheet["V4"].value == "Need to check"


def test_all_256_possible_sites_stay_within_excel_formula_argument_limit():
    report = report_fixture(())
    report["sites"] = list(range(256))
    _, book = workbook(report)
    sheet = book["raw"]
    assert sheet.max_column == 1801
    formulas = [rule.formula[0] for rules in sheet.conditional_formatting._cf_rules.values() for rule in rules]
    assert not any('COUNT(' in formula or 'COUNTIF(' in formula for formula in formulas)
