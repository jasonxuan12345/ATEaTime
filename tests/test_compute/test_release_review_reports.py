"""Saved population boundaries and complete review-report exports."""
from copy import deepcopy
from io import BytesIO
import json
import numpy as np
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from deepyield.compute.release_mmsa import build_release_mmsa_report
from deepyield.compute.release_site_diff_excel import render_site_diff_xlsx
from deepyield.compute.release_site_diff_pdf import render_site_diff_pdf
from deepyield.compute.release_chart_samples import ReportSampleResolver
from deepyield.compute.release_review_data import site_diff_rows, MATRIX_NOTE
from deepyield.compute.release_outliers import build_release_outlier_report
from deepyield.compute.release_outlier_excel import render_outlier_xlsx
from deepyield.compute.release_outlier_pdf import render_outlier_pdf
from tests.test_compute.test_release_mmsa import sites_frame
from deepyield.rules.loader import load_analysis_rules


def fixture(sites=2,parts=40):
    values={site:[1.+(i%9)*.01+(0.5 if site==0 and i==parts-1 else site*.08) for i in range(parts)] for site in range(sites)}
    frame,meta=sites_frame(values,0.,2.)
    frame['HEAD_NUM']=1;frame['LOT_ID']='LOT';frame['WAFER_ID']='W1';frame['STDF_PART_ID']=frame['__part_index'].astype(str)
    frame['DIE_X']=frame['__part_index']%10;frame['DIE_Y']=frame['__part_index']//10;frame['HBIN']=1
    frame.loc[frame.index[-1],['SBIN','HBIN']]=2
    frame.loc[frame.index[-1],'PART_FLAG']=8
    meta['test_catalog'][0]['name']='Long measurement name '*8
    return frame,meta


def site_fixture(sites=2,parts=40):
    frame,meta=fixture(sites,parts)
    report=build_release_mmsa_report(frame,meta,catalog=load_analysis_rules())
    return frame,meta,report


def outlier_fixture():
    frame,meta=fixture()
    report=build_release_outlier_report(frame,meta,minimum_samples_per_site=8)
    for b in report['baselines']:b['observation']={'preview':'Saved human note - keep this text.'}
    source={'dataset_id':'file-a','filename':'first.stdf','metadata':meta,'report':report}
    points=[]
    for b in report['baselines']:
        for r in frame.to_dict('records'):
            e=r['__test_records'][0]
            if r['SITE']==b['site'] and r['SBIN']==1:
                points.append({'part_index':r['__part_index'],'record_sequence':e['record_sequence'],'value':e['result'],
                    'site':r['SITE'],'baseline_id':b['baseline_id']})
    view={'points':points,'total_points':len(points),'sampled':False,'dies':[{'lot':'LOT','wafer':'W1','x':int(r.DIE_X),'y':int(r.DIE_Y)} for r in frame.itertuples()]}
    return {'sources':[source],'metadata':meta},lambda *args:view


def test_site_diff_excel_uses_saved_pass_fail_population_and_site_zero():
    frame,meta,report=site_fixture()
    report['items'][0]['diagnosis']={'auto_remark':'Automatic evidence.','remark':'Human saved remark.','knowledge_hash':'version-x'}
    original=deepcopy(report)
    data=render_site_diff_xlsx(report,metadata=meta)
    book=load_workbook(BytesIO(data['content']))
    assert book.sheetnames==['raw','Abnormal_Site_Diff']
    sheet=book['raw'];headers=[c.value for c in sheet[2]]
    assert 'S0_n' in headers and 'S0_Cg' not in headers
    assert sheet.cell(3,headers.index('Overall n')+1).value==len(frame)
    assert sheet.cell(3,headers.index('Overall Mean')+1).value==pytest.approx(report['items'][0]['overall_mean'])
    assert sheet.cell(3,headers.index('Site Std Ratio')+1).value==pytest.approx(report['items'][0]['var_ratio'])
    assert 'PASS/FAIL included' in sheet['A1'].value
    assert 'Human saved remark' in sheet.cell(3,len(headers)).value
    assert 'version-x' in sheet.cell(3,len(headers)).comment.text
    assert sheet.freeze_panes=='F3'
    assert original==report


@pytest.mark.parametrize('sites',[2,16])
def test_site_diff_pdf_is_vector_16x9_full_stats_and_annotations(sites):
    frame,meta,report=site_fixture(sites)
    item=report['items'][0]
    item.update(status='RISK',rule_id='MMSA-003',shift_severity='RED')
    item['diagnosis']={'remark':'Long saved remark. '*800}
    resolver=ReportSampleResolver(frame,meta['source_sha256'],identifiers=False)
    data=render_site_diff_pdf(report,metadata=meta,sample_provider=resolver,population='LOT / W1 / 25C')
    pdf=PdfReader(BytesIO(data['content']))
    assert len(pdf.pages)==2
    assert list(pdf.pages[0].mediabox)==[0,0,1152,648]
    text=pdf.pages[1].extract_text()
    for label in ('Test Sequence','Site Violin','Site Histogram','Site Statistics','S0'):assert label in text
    assert 'PASS/FAIL included' in pdf.pages[0].extract_text()
    assert any('Long saved remark' in a.get_object().get('/Contents','') for a in pdf.pages[1]['/Annots'])


def test_site_diff_unevaluated_does_not_become_pass_and_summary_only():
    _,meta,report=site_fixture()
    report['items'][0].update(status='NOT_EVALUATED',rule_id=None)
    data=site_diff_rows(report,meta)
    assert data['excluded']==1 and data['counts']['Pass']==0
    def forbidden(*a):pytest.fail('Excluded items must not resolve samples')
    pdf=render_site_diff_pdf(report,metadata=meta,sample_provider=forbidden,population='--')
    assert pdf['page_count']==1


def test_site_diff_5000_point_limit_does_not_limit_histogram_or_statistics():
    frame,meta,report=site_fixture(2,2600)
    report['items'][0].update(status='RISK',rule_id='MMSA-003')
    data=render_site_diff_pdf(report,metadata=meta,population='--',sample_provider=ReportSampleResolver(frame,meta['source_sha256']))
    text=PdfReader(BytesIO(data['content'])).pages[1].extract_text()
    assert 'display 5000 / 5200' in text
    assert report['items'][0]['n']==5200


def test_outlier_excel_preserves_references_candidates_hits_and_missing_cells():
    bundle,_=outlier_fixture();report=bundle['sources'][0]['report']
    assert report['chips']
    before=deepcopy(bundle)
    data=render_outlier_xlsx(bundle)
    book=load_workbook(BytesIO(data['content']))
    assert book.sheetnames==['Reference_Group_Stats','Candidate_Records','Candidate_Test_Matrix','Outlier_Hit_Details']
    refs=book.worksheets[0];headers=[c.value for c in refs[1]]
    assert refs.max_row==len(report['baselines'])+1
    assert any(refs.cell(i,headers.index('Site')+1).value==0 for i in range(2,refs.max_row+1))
    assert book['Outlier_Hit_Details'].max_row==report['summary']['candidate_hit_count']+1
    assert MATRIX_NOTE in book['Candidate_Test_Matrix']['A1'].comment.text
    assert 'Saved human note' in refs.cell(2,len(headers)).value
    assert bundle==before


def test_outlier_pdf_shows_both_main_views_and_does_not_merge_files():
    bundle,provider=outlier_fixture()
    other=deepcopy(bundle['sources'][0]);other['dataset_id']='file-b';other['filename']='second.stdf'
    bundle['sources'].append(other)
    output=render_outlier_pdf(bundle,sample_provider=provider)
    pdf=PdfReader(BytesIO(output['content']))
    text='\n'.join(p.extract_text() for p in pdf.pages)
    assert 'first.stdf' in text and 'second.stdf' in text
    detail=[p for p in pdf.pages if 'Reference Distribution' in p.extract_text()]
    assert len(detail)==2
    assert all('Test Sequence' in p.extract_text() for p in detail)
    assert any(MATRIX_NOTE in p.extract_text() for p in pdf.pages)
    assert output['candidate_count']==2
    assert len(pdf.pages)>=4


def test_outlier_empty_candidate_result_is_not_rendered_as_all_pass():
    bundle,provider=outlier_fixture()
    source=bundle['sources'][0]
    source['report']['chips']=[]
    for b in source['report']['baselines']:
        b.update(status='INSUFFICIENT_DATA',candidate_count=0)
    source['report']['summary']['unevaluated_baseline_count']=len(source['report']['baselines'])
    pdf=render_outlier_pdf(bundle,sample_provider=lambda *a:pytest.fail('No sample reads for empty candidates'))
    assert pdf['page_count']==1
    text=PdfReader(BytesIO(pdf['content'])).pages[0].extract_text()
    assert 'No candidates does not mean every group passed' in text


def test_site_diff_rejects_different_samples_even_with_same_count():
    frame,meta,report=site_fixture()
    report['items'][0].update(status='RISK',rule_id='MMSA-003')
    resolver=ReportSampleResolver(frame,meta['source_sha256'])
    def changed(item):return [{**p,'value':p['value']+1} for p in resolver(item)]
    with pytest.raises(ValueError,match='saved statistics'):
        render_site_diff_pdf(report,metadata=meta,population='--',sample_provider=changed)
