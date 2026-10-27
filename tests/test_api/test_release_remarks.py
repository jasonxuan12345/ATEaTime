from io import BytesIO
from copy import deepcopy
import openpyxl
from tests.test_api.test_release_check import client, upload, run
from app import release_check_service as service


def context(client):
    dataset = upload(client)
    result = run(client, dataset, modules=["cpk", "cgm", "mmsa"])
    root = f"/api/release-check/datasets/{dataset['id']}/runs/{result['id']}"
    return dataset, result, root


def test_auto_remarks_exist_before_manual_observations_and_agree_with_xlsx(client):
    _, result, root = context(client)
    observations = client.get(root + "/observations").json()
    assert observations["notes"] == {}
    sheet = openpyxl.load_workbook(BytesIO(client.get(root + "/observations.xlsx").content)).active
    rows = list(sheet.values)
    for module in ("cpk", "cgm", "mmsa"):
        cell = result["rows"][0][module]
        diagnosis = cell["diagnosis"]
        assert diagnosis["auto_remark"] == diagnosis["remark"]
        assert observations["remarks"][module + ":" + result["rows"][0]["key"]] == diagnosis
        values = next(values for values in rows[1:] if values[2] == module)
        assert (values[9] or '') == diagnosis["remark"]
    assert sheet.cell(2, 10).comment


def test_preview_save_json_excel_and_undo_share_one_canonical_remark(client):
    _, result, root = context(client)
    row = result["rows"][0]
    payload = dict(keys=[row["key"]], module="cgm", tags=["end_spikes"], suggestion="Check settling time",
                   remark="Review sample =1+1", state="待跟进", confirmed=True, overwrite=True, version=0)
    preview = client.post(root + "/observations/preview", json=payload)
    assert preview.status_code == 200
    saved = client.post(root + "/observations", json=payload).json()
    text = saved["notes"]["cgm:" + row["key"]]["remark"]
    assert text == preview.json()["remark"]
    assert "first/last sequence spikes observed" in text and "Review sample =1+1" in text
    report = client.get(root + "/report").json()
    assert report["result"]["rows"][0]["cgm"]["diagnosis"]["remark"] == text
    assert report["reports"]["cgm"]["items"][0]["diagnosis"]["remark"] == text
    assert report["result"]["rows"][0]["status"] == row["status"]
    sheet = openpyxl.load_workbook(BytesIO(client.get(root + "/observations.xlsx").content)).active
    values = next(values for values in list(sheet.values)[1:] if values[2] == "cgm")
    assert values[9] == text
    assert client.post(root + "/observations/preview", json=payload).status_code == 422
    undone = client.post(root + "/observations/undo/" + saved["undo_id"]).json()
    assert undone["notes"] == {}
    assert undone["remarks"]["cgm:" + row["key"]]["remark"] == row["cgm"]["diagnosis"]["auto_remark"]


def test_remarks_do_not_recompute_or_change_existing_metrics(client, monkeypatch):
    dataset, result, root = context(client)
    snapshot = deepcopy(result["rows"][0])
    def forbidden(*args, **kwargs):
        raise AssertionError("Original measurements must not be scanned for remarks")
    monkeypatch.setattr(service, "build_release_bin1_cpk_report", forbidden)
    monkeypatch.setattr(service, "build_release_cgm_report", forbidden)
    response = client.get(root + "/observations")
    assert response.status_code == 200
    assert client.get(root + "/report").json()["result"]["rows"][0] == snapshot


def test_outlier_remarks_use_own_snapshot_and_excel_matches(client):
    dataset = upload(client)
    response = client.post(f"/api/release-check/datasets/{dataset['id']}/outliers",
        json={"minimum_samples_per_site": 4})
    assert response.status_code == 200, response.text
    data = response.json()
    endpoint = f"/api/release-check/datasets/{dataset['id']}/outliers/{data['report_id']}.xlsx"
    response = client.get(endpoint)
    assert response.status_code == 200
    sheet = openpyxl.load_workbook(BytesIO(response.content)).active
    for values, baseline in zip(list(sheet.values)[1:], data["report"]["baselines"], strict=True):
        assert values[9] == baseline["diagnosis"]["remark"]
    assert client.get(endpoint.replace(data["report_id"], "expired")).status_code == 404


def test_engineer_can_edit_final_remark_without_erasing_automatic_evidence(client):
    _, result, root = context(client)
    row = result["rows"][0]
    automatic = row["cpk"]["diagnosis"]["auto_remark"]
    payload = dict(keys=[row["key"]], module="cpk", tags=[], confirmed=True, overwrite=True,
                   version=0, final_remark="Engineer reviewed the current samples; follow up on trim target.")
    preview = client.post(root + "/observations/preview", json=payload).json()
    saved = client.post(root + "/observations", json=payload).json()
    diagnosis = saved["remarks"]["cpk:" + row["key"]]
    assert diagnosis["remark"] == preview["remark"] == payload["final_remark"]
    assert diagnosis["auto_remark"] == automatic
    sheet = openpyxl.load_workbook(BytesIO(client.get(root + "/observations.xlsx").content)).active
    assert next(values for values in list(sheet.values)[1:] if values[2] == "cpk")[9] == payload["final_remark"]


def test_batch_cannot_copy_a_single_final_remark_to_other_tests(client):
    dataset, result, root = context(client)
    stored, _ = service.get_run(service.get_dataset(dataset["id"]), result["id"])
    other = deepcopy(stored["rows"][0])
    other["key"] = "another-test"
    other["number"] = 2002
    stored["rows"].append(other)
    payload = dict(keys=[stored["rows"][0]["key"], "another-test"], module="cpk", tags=[],
                   confirmed=True, overwrite=True, version=0, final_remark="Current S4 only.")
    response = client.post(root + "/observations", json=payload)
    assert response.status_code == 422
    assert "批量" in response.json()["detail"]
    assert client.get(root + "/observations").json()["notes"] == {}
