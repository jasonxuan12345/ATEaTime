"""Run association, source isolation, cache versions and report downloads."""
from io import BytesIO
from openpyxl import load_workbook
from pypdf import PdfReader
from app import release_check_service as release
from tests.test_api.test_release_outliers import client


def prepare(client):
    http,source=client
    response=http.post('/api/release-check/datasets/memory/run',json={'modules':['cpk','cgm','mmsa'],
        **{m:{'minimum_samples':2} for m in ('cpk','cgm','mmsa')}})
    assert response.status_code==200,response.text
    result=response.json()
    out=http.post('/api/release-check/datasets/memory/outliers',json={})
    assert out.status_code==200,out.text
    out=out.json()
    root=f"/api/release-check/datasets/memory/runs/{result['id']}"
    payload={'sources':[{'dataset_id':'memory','report_id':out['report_id']}],
             'observations':[{'dataset_id':'memory','group_id':out['report']['baselines'][0]['group_id'],
                              'note':{'remark':'Saved engineer comment','preview':'Saved engineer comment'}}]}
    response=http.put(root+'/outlier-reports',json=payload)
    assert response.status_code==200,response.text
    return http,source,root,payload,result


def test_four_exports_use_bound_memory_and_keep_user_notes(client):
    http,source,root,payload,result=prepare(client)
    original=source.runs[result['id']][1]['mmsa']['items'][0]['n']
    for module in ('site_diff','outlier'):
        for extension in ('xlsx','pdf'):
            response=http.get(root+f'/reports/{module}.{extension}')
            assert response.status_code==200,response.text[:500]
            assert ('SiteDiff_Report' if module=='site_diff' else 'Outlier_Report') in response.headers['content-disposition']
            if extension=='xlsx':
                book=load_workbook(BytesIO(response.content))
                assert len(book.sheetnames)==(2 if module=='site_diff' else 4)
                if module=='outlier':assert any('Saved engineer comment' in str(c.value) for row in book.worksheets[0] for c in row)
            else:assert PdfReader(BytesIO(response.content)).pages
    assert source.runs[result['id']][1]['mmsa']['items'][0]['n']==original
    assert http.get(root+'/report').status_code==200  # Association is private, not part of result JSON.


def test_manifest_cache_reuses_same_snapshot_and_invalidates_saved_notes(client):
    http,source,root,payload,result=prepare(client)
    path=root+'/reports/outlier.xlsx?delivery=file&lang=en'
    first=http.get(path);assert first.status_code==200,first.text
    ready=first.json();download=http.get(ready['url'])
    assert download.content.startswith(b'PK')
    again=http.get(path).json()
    assert http.get(again['url']).content==download.content
    count=len(source.report_files.entries)
    assert http.put(root+'/outlier-reports',json=payload).json()['version']==1
    payload['observations'][0]['note']={'preview':'Updated saved note'}
    assert http.put(root+'/outlier-reports',json=payload).json()['version']==2
    second=http.get(path).json();new=http.get(second['url']).content
    assert new!=download.content
    book=load_workbook(BytesIO(new))
    assert any('Updated saved note' in str(c.value) for row in book.worksheets[0] for c in row)
    assert len(source.report_files.entries)>count


def test_rejects_unbound_outlier_and_unrelated_observations(client):
    http,source,root,payload,result=prepare(client)
    payload['observations'][0]['group_id']='unrelated'
    assert http.put(root+'/outlier-reports',json=payload).status_code==422
    assert source.outlier_report_exports[result['id']]['version']==1
    source.outlier_report_exports.clear()
    assert http.get(root+'/reports/outlier.pdf').status_code==404


def test_expired_source_and_pre_cancelled_report_are_not_exported(client):
    http,source,root,payload,result=prepare(client)
    http.delete('/api/release-check/jobs/cancel-report')
    assert http.get(root+'/reports/outlier.pdf?job_id=cancel-report').status_code==422
    assert not source.report_files.entries
    saved=source.outlier_report_exports[result['id']]
    saved['sources'][0]['dataset_id']='removed-source'
    assert http.get(root+'/reports/outlier.xlsx').status_code==404


def test_run_keeps_exact_outlier_snapshot_after_recent_list_eviction(client):
    http,source,root,payload,result=prepare(client)
    source.outlier_runs.clear()
    response=http.get(root+'/reports/outlier.pdf')
    assert response.status_code==200,response.text[:200]
    assert 'Saved engineer comment' in '\n'.join(p.extract_text() for p in PdfReader(BytesIO(response.content)).pages)
    report_id=payload['sources'][0]['report_id']
    full=http.get(f'/api/release-check/datasets/memory/outliers/{report_id}/report')
    assert full.status_code==200,full.text
    baseline=full.json()['baselines'][0]
    evidence=http.get(f'/api/release-check/datasets/memory/outliers/{report_id}/baselines/{baseline["baseline_id"]}/diagnosis')
    assert evidence.status_code==200 and evidence.json()==baseline['diagnosis']
