"""File downloads, concurrent cache misses, invalidation and bounded storage."""
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from threading import Event

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
import pytest

from app.release_export_cache import ReportFiles, register_download, take_download
from app import release_chart_notes as notes
from app.release_pdf_service import export_run_cgm_pdf
from app.release_excel_service import export_run_xlsx
from app.routers.release_check import router
from tests.test_api.test_release_excel import sources, run_pair
from tests.test_api.test_release_outliers import client


def test_cache_combines_concurrent_requests_and_keeps_download_alive_after_eviction():
    cache = ReportFiles(budget=10)
    started, finish = Event(), Event()
    calls = []
    def write(output):
        calls.append(1)
        started.set()
        assert finish.wait(5)
        output.write(b'12345')
        return {'filename': 'file.xlsx'}
    with ThreadPoolExecutor(max_workers=2) as pool:
        one = pool.submit(cache.build, 'same', write, file_delivery=True)
        assert started.wait(5)
        two = pool.submit(cache.build, 'same', write, file_delivery=True)
        finish.set()
        first, second = one.result(), two.result()
    assert len(calls) == 1 and first['_file'] is second['_file']
    ready = register_download(first['_file'])
    token = ready['url'].rsplit('/', 1)[1]
    for index in range(3):
        cache.build(index, write)
    assert len(cache.entries) == 2
    assert take_download(token).path.read_bytes() == b'12345'
    assert take_download(token).path.read_bytes() == b'12345'  # Native range/retry downloads.


def test_failed_export_does_not_enter_cache():
    cache = ReportFiles()
    def fail(output):
        output.write(b'partial')
        raise ValueError('failed')
    with pytest.raises(ValueError, match='failed'):
        cache.build('broken', fail)
    assert not cache.entries


def test_all_export_manifests_download_actual_files_and_reuse_cache(sources, monkeypatch):
    result = run_pair()
    app = FastAPI()
    app.include_router(router)
    with TestClient(app) as client:
        base = f'/api/release-check/datasets/a/runs/{result["id"]}'
        for suffix in ('reports/cpk.xlsx', 'reports/cgm.xlsx', 'reports/cpk.pdf', 'reports/cgm.pdf', 'observations.xlsx'):
            ready = client.get(base + '/' + suffix + '?delivery=file&job_id=export-file')
            assert ready.status_code == 200, ready.text
            manifest = ready.json()
            downloaded = client.get(manifest['url'])
            assert downloaded.status_code == 200
            assert len(downloaded.content) == manifest['size']
            assert downloaded.content.startswith(b'%PDF' if suffix.endswith('.pdf') else b'PK')
            assert downloaded.headers['cache-control'] == 'no-store'
            if suffix == 'observations.xlsx':
                assert load_workbook(BytesIO(downloaded.content)).active.max_row > 1
        assert client.get('/api/release-check/report-downloads/unknown').status_code == 404
    import app.release_pdf_service as pdf
    monkeypatch.setattr(pdf, 'render_cgm_pdf', lambda *a, **k: pytest.fail('must reuse PDF'))
    assert export_run_cgm_pdf('a', result['id'])['content'].startswith(b'%PDF')


def test_module_specific_cache_invalidation_and_observation_copy_scope(sources, monkeypatch):
    owner, _ = sources
    result = run_pair()
    first = export_run_xlsx('a', result['id'], 'cpk')
    row = next(row for row in result['rows'] if row.get('cgm'))
    notes.save_observations('a', result['id'], dict(keys=[row['key']], module='cgm', tags=[],
        suggestion='', remark='', final_remark='Updated CGM only', state='待跟进', confirmed=True, overwrite=True, version=0))
    assert export_run_xlsx('a', result['id'], 'cpk') is first
    class NoCopy:
        def __deepcopy__(self, memo):
            pytest.fail('unused history and site details must not be copied')
    result['observation_history'].append(NoCopy())
    row['cgm']['sites'] = NoCopy()
    output = notes.export_observations_xlsx('a', result['id'])
    sheet = load_workbook(BytesIO(output)).active
    assert any(cell.value == 'Updated CGM only' for cell in sheet['J'])


def test_cached_pdf_tracks_saved_remark_and_undo(sources, monkeypatch):
    from datetime import datetime, timezone
    from functools import partial
    from app import release_pdf_service as pdf
    # Compare report content across undo, independent of elapsed rendering time.
    monkeypatch.setattr(pdf, 'render_cgm_pdf', partial(
        pdf.render_cgm_pdf, generated_at=datetime(2026, 1, 1, tzinfo=timezone.utc)))
    owner, _ = sources
    result = run_pair()
    owner.runs[result['id']][1]['cgm']['items'][0]['supported_by_bin1_cpk'] = None
    first = export_run_cgm_pdf('a', result['id'])
    row = next(row for row in result['rows'] if row.get('cgm'))
    saved = notes.save_observations('a', result['id'], dict(keys=[row['key']], module='cgm', tags=[],
        suggestion='', remark='', final_remark='Cached PDF updated conclusion', state='待跟进',
        confirmed=True, overwrite=True, version=0))
    updated = export_run_cgm_pdf('a', result['id'])
    assert 'Cached PDF updated conclusion' in PdfReader(BytesIO(updated['content'])).pages[1].extract_text()
    notes.undo_observations('a', result['id'], saved['undo_id'])
    restored = export_run_cgm_pdf('a', result['id'])
    assert PdfReader(BytesIO(restored['content'])).pages[1].extract_text() == PdfReader(BytesIO(first['content'])).pages[1].extract_text()


def test_outlier_file_download(client):
    http, _ = client
    run = http.post('/api/release-check/datasets/memory/outliers', json={}).json()
    ready = http.get(f'/api/release-check/datasets/memory/outliers/{run["report_id"]}.xlsx?delivery=file').json()
    response = http.get(ready['url'])
    assert response.status_code == 200 and len(response.content) == ready['size']
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet.title == 'Outlier diagnostics' and sheet.max_row > 1

