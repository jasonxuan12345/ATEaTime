"""HTTP attachment download and compatibility with existing observation exports."""
from io import BytesIO
from urllib.parse import unquote

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.routers.release_check import router
from tests.test_api.test_release_excel import sources, run_pair


def test_real_workbook_http_downloads_and_existing_observations_route(sources):
    result = run_pair()
    app = FastAPI()
    app.include_router(router)
    base = f"/api/release-check/datasets/a/runs/{result['id']}"
    with TestClient(app) as client:
        for module, sheets, filename in (
            ('cpk', ['raw', 'low_CPK_Item', 'CPK_too_Good', 'Test_without_unit'], '_Device_a_25_Bin1_CPK_Report.xlsx'),
            ('cgm', ['raw'], '_Device_b_Cgm_Report.xlsx')):
            response = client.get(f'{base}/reports/{module}.xlsx')
            assert response.status_code == 200, response.text
            assert response.headers['content-type'].endswith('spreadsheetml.sheet')
            assert filename in unquote(response.headers['content-disposition'])
            assert load_workbook(BytesIO(response.content)).sheetnames == sheets
        assert client.get(base + '/observations.xlsx').status_code == 200
        assert client.get(base + '/reports/mmsa.xlsx').status_code == 422
        assert client.get(base.replace(result['id'], 'expired') + '/reports/cpk.xlsx').status_code == 404
