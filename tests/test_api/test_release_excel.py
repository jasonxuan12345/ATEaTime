"""Run export contracts without a network server: sources, snapshots and responses."""
from copy import deepcopy
from io import BytesIO
import json
from types import SimpleNamespace
from urllib.parse import unquote

from fastapi import HTTPException
from openpyxl import load_workbook
import pytest

from app import release_check_service as service
from app import release_chart_notes as notes
from app.release_excel_service import export_run_xlsx
from app.routers.release_check import capability_xlsx, router
from tests.test_compute.test_release_cpk import dataset


def add_source(tmp_path, key, values, low=1, high=15):
    frame, metadata = dataset(values, low, high)
    frame['__conditions_json'] = '{}'
    metadata.update(source_file=key + '.stdf', source_sha256=key * 64,
        stdf_metadata={'mir': {'part_type': 'Device_' + key, 'test_temperature': '25' if key == 'a' else '125'}})
    source = service.ReleaseDataset(SimpleNamespace(name=str(tmp_path)), tmp_path / (key + '.parquet'),
                                    key + '.stdf', frame, metadata)
    service._datasets[key] = source
    return source


@pytest.fixture
def sources(tmp_path, monkeypatch):
    monkeypatch.setattr(service, '_datasets', {})
    a = add_source(tmp_path, 'a', (7.9, 7.9, 8, 8, 8.1, 8.1))
    b = add_source(tmp_path, 'b', (1.1, 1.1, 8, 8, 14.9, 14.9))
    return a, b


def run_pair(**selections):
    return service.run_checks('a', {'cpk': selections.get('cpk', {}), 'cgm': selections.get('cgm', {})},
                              ['cpk', 'cgm'], {'cpk': 'a', 'cgm': 'b'})


def test_independent_sources_selections_cpk_support_and_real_filenames(sources):
    a, _ = sources
    result = run_pair(cgm={'sites': [1], 'part_indices': [1, 3, 5]})
    _, reports = a.runs[result['id']]
    cpk = export_run_xlsx('a', result['id'], 'cpk')
    cgm = export_run_xlsx('a', result['id'], 'cgm')
    assert '_Device_a_25_Bin1_CPK_Report.xlsx' in cpk['filename']
    assert '_Device_b_Cgm_Report.xlsx' in cgm['filename']
    cp = load_workbook(BytesIO(cpk['content']))['raw']
    cg = load_workbook(BytesIO(cgm['content']))['raw']
    assert cp['E2'].value == 6 and cp['I2'].value == 8
    assert cp['L2'].value == pytest.approx(reports['cpk']['items'][0]['cpk'])
    assert cg.max_column == 16 and cg['F2'].value == 1
    assert cg['K4'].value == 3 and cg['F4'].value == 8
    assert cg['G4'].value == pytest.approx(6.9)
    assert cg['N4'].value == cp['L2'].value and cg['O4'].value == 'Pass'
    assert json.loads(load_workbook(BytesIO(cgm['content'])).properties.description)['bin1_cpk_source']['source_file'] == 'a.stdf'


def test_export_does_not_scan_recompute_or_copy_measurement_references(sources, monkeypatch):
    a, b = sources
    result = run_pair()
    _, reports = a.runs[result['id']]
    class Forbidden:
        def __iter__(self):
            raise AssertionError('export must not iterate measurements')
        def __deepcopy__(self, memo):
            raise AssertionError('export must not copy measurements')
    def forbidden(*args, **kwargs):
        raise AssertionError('export must not calculate statistics')
    for report in reports.values():
        for item in report['items']:
            item['sample_refs'] = Forbidden()
    a.frame = b.frame = Forbidden()
    monkeypatch.setattr(service, 'build_release_bin1_cpk_report', forbidden)
    monkeypatch.setattr(service, 'build_release_cgm_report', forbidden)
    for module in ('cpk', 'cgm'):
        assert export_run_xlsx('a', result['id'], module)['content'].startswith(b'PK')


def test_unexecuted_cpk_uses_its_source_catalog(sources):
    a, b = sources
    for source, unit, low in ((a, 'A', 100), (b, 'Hz', 0)):
        source.metadata['test_catalog'].append(dict(column='NEVER', test_num=9999,
            name='NOT_RUN', unit=unit, low_limit=low, high_limit=low + 10))
    result = run_pair()
    sheet = load_workbook(BytesIO(export_run_xlsx('a', result['id'], 'cpk')['content']))['raw']
    assert (sheet['D3'].value, sheet['G3'].value, sheet['H3'].value) == ('A', 100, 110)


def test_manual_final_cpk_remark_and_undo_are_current(sources, tmp_path):
    add_source(tmp_path, 'a', (4, 5, 6), low=1, high=9)
    result = run_pair()
    row = next(row for row in result['rows'] if row.get('cpk'))
    before = deepcopy(result)
    first = export_run_xlsx('a', result['id'], 'cpk')
    saved = notes.save_observations('a', result['id'], dict(keys=[row['key']], module='cpk',
        tags=[], suggestion='', remark='', final_remark='Engineer saved final', state='待跟进',
        confirmed=True, overwrite=True, version=0))
    updated = export_run_xlsx('a', result['id'], 'cpk')
    assert load_workbook(BytesIO(updated['content']))['low_CPK_Item']['M2'].value == 'Engineer saved final'
    notes.undo_observations('a', result['id'], saved['undo_id'])
    undone = export_run_xlsx('a', result['id'], 'cpk')
    assert load_workbook(BytesIO(undone['content']))['low_CPK_Item']['M2'].value == load_workbook(BytesIO(first['content']))['low_CPK_Item']['M2'].value
    assert before['rows'][0]['cpk']['value'] == result['rows'][0]['cpk']['value']


def test_cgm_remark_matches_saved_page_and_updates_after_save_and_undo(sources):
    result = run_pair()
    row = next(row for row in result['rows'] if row.get('cgm'))
    automatic = row['cgm']['diagnosis']['remark']
    first = export_run_xlsx('a', result['id'], 'cgm')
    sheet = load_workbook(BytesIO(first['content']))['raw']
    assert sheet.cell(3, sheet.max_column).value == 'Remark'
    assert sheet.cell(4, sheet.max_column).value == automatic
    saved = notes.save_observations('a', result['id'], dict(keys=[row['key']], module='cgm',
        tags=[], suggestion='', remark='', final_remark='CGM engineer final', state='待跟进',
        confirmed=True, overwrite=True, version=0))
    updated = load_workbook(BytesIO(export_run_xlsx('a', result['id'], 'cgm')['content']))['raw']
    assert updated.cell(4, updated.max_column).value == 'CGM engineer final'
    notes.undo_observations('a', result['id'], saved['undo_id'])
    undone = load_workbook(BytesIO(export_run_xlsx('a', result['id'], 'cgm')['content']))['raw']
    assert undone.cell(4, undone.max_column).value == automatic


def test_response_headers_route_and_expired_source_or_run(sources):
    result = run_pair()
    response = capability_xlsx('a', result['id'], 'cgm')
    assert response.media_type.endswith('spreadsheetml.sheet')
    assert response.headers['Cache-Control'] == 'no-store'
    assert int(response.headers['Content-Length']) == len(response.body)
    assert '_Device_b_Cgm_Report.xlsx' in unquote(response.headers['Content-Disposition'])
    assert 'Server-Timing' in response.headers
    paths = [route.path for route in router.routes]
    assert '/api/release-check/datasets/{dataset_id}/runs/{run_id}/reports/{module}.xlsx' in paths
    assert '/api/release-check/datasets/{dataset_id}/runs/{run_id}/observations.xlsx' in paths
    for dataset_id, run_id, module in [('b', result['id'], 'cpk'), ('a', 'expired', 'cpk'),
                                        ('a', result['id'], 'mmsa')]:
        with pytest.raises(HTTPException) as error:
            capability_xlsx(dataset_id, run_id, module)
        assert error.value.status_code in (404, 422)
    service.remove_dataset('b')
    with pytest.raises(HTTPException) as error:
        capability_xlsx('a', result['id'], 'cgm')
    assert error.value.status_code == 404


def test_source_mismatch_and_missing_module_are_rejected(sources):
    a, b = sources
    result = run_pair()
    b.metadata['source_sha256'] = 'changed'
    with pytest.raises(ValueError, match='来源文件不一致'):
        export_run_xlsx('a', result['id'], 'cgm')
    del a.runs[result['id']][1]['cpk']
    with pytest.raises(KeyError, match='没有这个模块'):
        export_run_xlsx('a', result['id'], 'cpk')


def test_repeated_download_uses_cache_and_removed_sources_still_fail(sources, monkeypatch):
    import app.release_excel_service as exporter
    result = run_pair()
    first = export_run_xlsx('a', result['id'], 'cgm')
    monkeypatch.setattr(exporter, 'render_cgm_xlsx', lambda *args, **kwargs: pytest.fail('cached download must not rebuild Excel'))
    assert export_run_xlsx('a', result['id'], 'cgm') is first
    service.remove_dataset('b')
    with pytest.raises(KeyError):
        export_run_xlsx('a', result['id'], 'cgm')


def test_cache_budget_and_run_eviction(sources, monkeypatch):
    import app.release_excel_service as exporter
    a, _ = sources
    result = run_pair()
    monkeypatch.setattr(exporter, 'MAX_CACHE_BYTES', 1)
    export_run_xlsx('a', result['id'], 'cpk')
    assert not a.excel_cache
    monkeypatch.setattr(exporter, 'MAX_CACHE_BYTES', 16 * 1024 ** 2)
    export_run_xlsx('a', result['id'], 'cpk')
    assert len(a.excel_cache) == 1
    for _ in range(3):
        run_pair()
    assert result['id'] not in a.runs and not a.excel_cache
